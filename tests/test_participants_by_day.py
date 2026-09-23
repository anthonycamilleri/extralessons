"""The class list's "Download participants by day" workbook: every enrolled
child of every class this term that is not cancelled, one sheet per weekday."""
import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.catalog.models import ActivityClass
from apps.enrollments.models import Enrollment

from .factories import (
    ActivityClassFactory,
    AdminFactory,
    ChildFactory,
    EnrollmentFactory,
    ReadOnlyAdminFactory,
    SuperAdminFactory,
    TermFactory,
)

pytestmark = pytest.mark.django_db

URL = reverse("admin:catalog_activityclass_participants_by_day")
S = Enrollment.Status


def enrol(cls, first_name, status=S.ENROLLED):
    return EnrollmentFactory(
        activity_class=cls, child=ChildFactory(first_name=first_name), status=status
    )


def download(client, user):
    client.force_login(user)
    response = client.get(URL)
    assert response.status_code == 200
    assert response["Content-Disposition"].endswith('"participants-by-day.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    return {
        sheet.title: [row for row in sheet.iter_rows(values_only=True)]
        for sheet in workbook.worksheets
    }


def test_one_sheet_per_day_with_student_and_class(client):
    chess = ActivityClassFactory(title="Chess", weekday=0, location="Library")
    judo = ActivityClassFactory(title="Judo", weekday=2)
    enrol(chess, "Ada")
    enrol(judo, "Bea")
    enrol(judo, "Cal", status=S.WAITLISTED)
    enrol(judo, "Dot", status=S.CANCELLED)

    sheets = download(client, SuperAdminFactory())

    assert list(sheets) == ["Monday", "Wednesday"]
    header, *rows = sheets["Monday"]
    assert header[:3] == ("Student", "School class", "Class")
    assert rows == [("Ada Tester", "P3E", "Chess", "15:00–16:00", "Library", chess.provider.name)]
    assert [row[0] for row in sheets["Wednesday"][1:]] == ["Bea Tester"]


def test_cancelled_classes_and_inactive_terms_are_left_out(client):
    enrol(ActivityClassFactory(weekday=1, status=ActivityClass.Status.CANCELLED), "Eve")
    enrol(ActivityClassFactory(weekday=3, term=TermFactory(is_active=False)), "Fay")
    enrol(ActivityClassFactory(weekday=4), "Gus")

    sheets = download(client, SuperAdminFactory())

    assert list(sheets) == ["Friday"]


def test_a_regular_admin_gets_only_their_classes(client):
    admin = AdminFactory()
    mine = ActivityClassFactory(weekday=0)
    mine.administrators.add(admin)
    enrol(mine, "Hal")
    enrol(ActivityClassFactory(weekday=0), "Ivy")

    sheets = download(client, admin)

    assert [row[0] for row in sheets["Monday"][1:]] == ["Hal Tester"]


def test_read_only_admin_can_download_and_empty_workbook_still_opens(client):
    assert list(download(client, ReadOnlyAdminFactory())) == ["No participants"]


def test_button_on_the_class_list(client):
    client.force_login(SuperAdminFactory())
    page = client.get(reverse("admin:catalog_activityclass_changelist")).content.decode()
    assert URL in page
